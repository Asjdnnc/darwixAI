"""Build a concise, submission-ready Word report from the approved Markdown plan."""
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = Path("output/AI_Engineer_Assessment_Report.docx")
NAVY = "17365D"
PALE = "EAF1F8"
GRAY = "D9D9D9"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell, color=GRAY):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = OxmlElement(f"w:{edge}")
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "6")
        tag.set(qn("w:color"), color)
        borders.append(tag)
    tc_pr.append(borders)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = header
        shade(cell, NAVY)
        set_cell_border(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
    for row_i, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
            set_cell_border(cells[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_i % 2:
                shade(cells[i], PALE)
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    doc.add_paragraph()
    return table


def add_bullets(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def build():
    OUT.parent.mkdir(exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(11)
    for name, size in [("Title", 24), ("Heading 1", 16), ("Heading 2", 13)]:
        styles[name].font.name = "Aptos Display"
        styles[name].font.size = Pt(size)
        styles[name].font.color.rgb = RGBColor(0, 0, 0)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("AI Engineer Assessment Report")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Groq Browser Based Prototype Plan").italic = True
    intro = doc.add_paragraph()
    intro.add_run("Purpose. ").bold = True
    intro.add_run("This report explains the integrated implementation plan, evidence requirements, test strategy, and delivery order for the four-part AI Engineer Assessment. The proposed solution uses the Groq API for reasoning, speech-to-text, English speech synthesis, and call analysis.")

    doc.add_heading("Assignment Interpretation", level=1)
    add_table(doc, ["Question", "Required outcome", "Submission evidence"], [
        ["1 Voice agent", "Grounded lead-qualification voice flow connected to the knowledge base", "Three recorded calls, transcripts, outcomes, safe fallback and escalation"],
        ["2 Knowledge base", "Clean, traceable, searchable business knowledge", "Schema, source tracking, citations and five retrieval evaluations"],
        ["3 Localized bots", "Philippines and Indonesia financial voice prototypes", "Two recorded calls per market, localization and ASR/TTS observations"],
        ["4 Live nudges", "Real-time call analysis and actionable agent assist", "Live/replay demo, latency report, false-positive controls"],
    ], [1.2, 3.2, 2.4])

    doc.add_heading("System Design", level=1)
    doc.add_paragraph("A browser microphone records short audio chunks and sends them to FastAPI. The service uses Groq Whisper for transcription, retrieves traceable knowledge chunks, sends only relevant context to Groq for a grounded answer, and returns citations. Incoming transcript chunks are analyzed immediately for signals and then passed through a nudge controller that applies confidence thresholds, deduplication, and cooldowns.")
    add_table(doc, ["Component", "Selected technology", "Reason"], [
        ["Voice channel", "Browser microphone plus Groq", "No telephony dependency; Groq Whisper transcription and English TTS"],
        ["Reasoning", "Groq", "Grounded generation and structured classification with provider configuration kept in environment variables"],
        ["Application", "Python FastAPI", "Typed API contract, clear OpenAPI documentation, simple webhook hosting"],
        ["Knowledge base", "Traceable chunks with source metadata", "Supports auditability, citations, source versioning, and safe answers"],
        ["Real-time controls", "Rules plus Groq extension", "Reproducible baseline and clear path to model-assisted signal detection"],
    ], [1.3, 1.8, 3.7])

    doc.add_heading("Question 1 Knowledge Grounded Voice Agent", level=1)
    doc.add_paragraph("The selected flow is health-insurance lead qualification. The agent collects customer details, requested coverage, location, preferred callback time, and consent. It calls the retrieval tool before answering policy, product, qualification, pricing, or payment-support questions. When sources are absent or inconclusive, it says that verified information is unavailable and offers human assistance.")
    add_bullets(doc, ["Required tests: cooperative customer, objection, conflicting details, unsupported request, and human escalation.", "Optional business action: create a synthetic lead or callback request after consent.", "Evidence: recording, transcript, citations used, qualification result, and escalation/lead outcome."])

    doc.add_heading("Question 2 Production Ready Knowledge Base", level=1)
    doc.add_paragraph("The data pipeline stages web pages, PDFs, forms, policy documents, and product material; removes boilerplate and duplicates; standardizes headings and terminology; flags parsing failures; and detects/redacts PII before indexing. Every source receives a stable identifier, version, category, source reference, and PII flag. Chunks retain their parent source metadata.")
    add_table(doc, ["Field", "Example"], [
        ["record_id", "policy-disclosure"], ["category", "policy"], ["version", "1.0"], ["source_ref", "sample/policies/disclosure-v1.md"], ["pii_present", "false"], ["chunk_id", "policy-disclosure:0"],
    ], [2.0, 4.8])
    doc.add_paragraph("The prototype starts with transparent lexical ranking for reproducible testing. Production should use embeddings, vector search, and reranking while preserving the same metadata and citation contract.")

    doc.add_heading("Question 3 Native Language Voice Bots", level=1)
    add_table(doc, ["Market", "Prototype", "Localization requirements"], [
        ["Philippines", "Life insurance or bancassurance", "English, Filipino/Tagalog, Taglish; premium, policy, beneficiary, rider, lapse, coverage, bank referral"],
        ["Indonesia", "Consumer-finance reminder or follow-up", "Formal/colloquial Bahasa, finance loanwords, regional-accent sample; cicilan, tenor, denda, DP, jatuh tempo, angsuran, pembiayaan"],
    ], [1.2, 2.0, 3.6])
    doc.add_paragraph("The localized scripts must not be literal translations. Dates, currency explanations, politeness, objections, consent, and escalation are adapted for each market. Every fallback remains in the speaker's language. Groq currently documents English and Saudi Arabic TTS voices only, so the Philippines and Indonesia prototypes return localized text and document the native-TTS limitation.")

    doc.add_heading("Question 4 Live Insights and Nudges", level=1)
    doc.add_paragraph("Transcript chunks are processed as they arrive during a live call or a recording replayed at real-time speed. The system detects missed cross-sell opportunities, compliance gaps, frustration, payment difficulty, and callback needs. Nudge messages are short, actionable, and expire after 90 seconds.")
    add_table(doc, ["Signal", "Nudge", "Control"], [
        ["Missing disclosure", "Confirm the required disclosure before proceeding", "High-confidence detection and topic cooldown"],
        ["Second vehicle", "Ask about a multi-vehicle quote", "One nudge per active topic"],
        ["Rising frustration", "Acknowledge concern and clarify the next step", "Avoid repeat alerts"],
        ["Payment difficulty", "Offer approved support or callback path", "Never promise an unapproved exception"],
    ], [1.5, 3.2, 2.1])

    doc.add_heading("Testing and Acceptance", level=1)
    add_bullets(doc, ["Run automated tests for retrieval traceability, safe unsupported-question handling, nudge generation, and duplicate suppression.", "Evaluate at least five retrieval queries across product, policy, qualification, FAQ, and objection categories. Record retrieved chunks, citations, explanation, and verdict.", "Record at least three Question 1 calls and two Question 3 calls per market.", "Measure P50/P95 overall and component latency using event timestamps from receipt through nudge display.", "Run ambiguous/noisy input tests to estimate false positives and demonstrate alert suppression."])

    doc.add_heading("Delivery Checklist", level=1)
    add_bullets(doc, ["Repository with README, setup instructions, source code, test fixtures, and .env.example.", "Architecture, knowledge-base schema, Groq configuration, browser workflow, test plan, demo runbook, transcripts, recordings, and latency results.", "Video walkthrough: system overview, Q1-Q2 connection, multilingual behavior, live nudge generation, fallbacks, limitations, and production plan.", "No API keys, secrets, recordings with real customer information, or other sensitive data committed to source control."])

    doc.add_heading("Known Limitations and Production Plan", level=1)
    doc.add_paragraph("The starter repository uses local in-memory knowledge and deterministic baseline signals so that it remains understandable and testable without provider credentials. Before production, replace local state with PostgreSQL and vector indexing, add Groq structured classification evaluated on labeled data, implement data retention and encryption, add browser-session security and observability, queue work for scale, and obtain native-language/compliance review.")

    doc.save(OUT)


if __name__ == "__main__":
    build()
